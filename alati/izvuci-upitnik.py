# -*- coding: utf-8 -*-
"""
Izvlaci katalog IT upitnika (dijelovi A, B, C) iz izvornog Excel fajla u JSON.

Izvorni fajl je prazan sablon - ne sadrzi nijedan popunjen odgovor, pa se
prenose samo struktura, tekstovi pitanja, uputstva i liste ponudjenih odgovora.

Struktura izvornog sheeta:
    kolona A  oznaka dijela (A / B / C) na redu zaglavlja dijela
    kolona B  kod sekcije (A1, B1, C1...) ili naziv dijela
    kolona C  redni broj pitanja (1, 15b, 43c) ili kod podoblasti (C1.1)
    kolona D  tekst pitanja, odnosno naziv podoblasti
    kolona E  odgovor            (prazno u sablonu)
    kolona F  objasnjenje        (dio C: obrazlozenje prednosti)
    kolona G  uputstvo (A i B)   (dio C: obrazlozenje slabosti)
    kolona H  ocjena             (dio B: ocjena rizika, dio C: ocjena zrelosti)

Pokretanje:
    python alati/izvuci-upitnik.py "putanja/do/3.Samoprocjena nivoa IT rizika.xlsx"

Izlaz: prisma/podaci/upitnik.json
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

IZLAZ = Path(__file__).resolve().parent.parent / 'prisma' / 'podaci' / 'upitnik.json'

PRVI_RED = 14
POSLJEDNJI_RED_PITANJA = 480
PRVI_RED_REZULTATA = 485
POSLJEDNJI_RED_REZULTATA = 534

RE_SEKCIJA = re.compile(r'^([ABC]\d+)$')
RE_PODOBLAST = re.compile(r'^(C\d+\.\d+)\.?$')
RE_BROJ_PITANJA = re.compile(r'^\d+[a-z]?$')

# Brojcano pitanje se prepoznaje po upitnoj rijeci BILO GDJE u tekstu, jer
# veliki broj podpitanja pocinje kvalifikatorom ("Od tog iznosa koliko...",
# "Od toga, koji je broj...", "Po vasoj procjeni koliki je...").
RE_BROJ = re.compile(
    r'(\bkolik\w*\b'
    r'|\bkoji je (ukupan )?broj\b'
    r'|\bkoji je procenat\b'
    r'|\bu koliko slučajeva\b'
    r'|\bprosječno vrijeme\b'
    r'|^broj\s)', re.I)
RE_DA_NE = re.compile(r'^(da li|postoji li|jesu li|je li)', re.I)


def ocisti(v):
    return '' if v is None else re.sub(r'\s+', ' ', str(v)).strip()


def mapa_validacija(ws):
    """Mapira adresu celije (npr. E30) na listu ponudjenih odgovora."""
    mapa = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != 'list' or not dv.formula1:
            continue
        sirovo = str(dv.formula1).strip().strip('"')
        opcije = [o.strip() for o in sirovo.split(',') if o.strip()]
        if not opcije:
            continue
        for opseg in dv.sqref.ranges:
            for red in ws.iter_rows(
                min_row=opseg.min_row, max_row=opseg.max_row,
                min_col=opseg.min_col, max_col=opseg.max_col,
            ):
                for celija in red:
                    mapa[celija.coordinate] = opcije
    return mapa


def kanonski_nazivi(ws):
    """Zvanicni nazivi sekcija iz bloka Rezultat samoprocjene."""
    nazivi = {}
    for r in range(PRVI_RED_REZULTATA, POSLJEDNJI_RED_REZULTATA + 1):
        oznaka = ocisti(ws.cell(r, 5).value)
        if not oznaka:
            continue
        m = re.match(r'^([ABC]\d+(?:\.\d+)?)\.?\s+(.*)$', oznaka)
        if not m:
            continue
        kod, naziv = m.group(1), m.group(2)
        naziv = re.sub(r'\s*-\s*zrelost$', '', naziv).strip()
        if naziv.lower().startswith('indikatori it rizika ukupno'):
            continue
        nazivi[kod] = naziv
    return nazivi


def tip_odgovora(dio, tekst, adresa_e, validacije):
    if dio == 'C':
        return 'DA_NE_DJELIMICNO', []

    opcije = validacije.get(adresa_e)
    if opcije:
        if [o.lower() for o in opcije] == ['da', 'ne']:
            return 'DA_NE', []
        return 'IZBOR', opcije

    if RE_DA_NE.match(tekst):
        return 'DA_NE', []
    if RE_BROJ.search(tekst):
        return 'BROJ', []
    return 'TEKST', []


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=False)
    ws = wb['IT upitnik']

    validacije = mapa_validacija(ws)
    nazivi = kanonski_nazivi(ws)

    sekcije = []
    dio = None
    oblast = None      # tekuca oblast dijela C (C1...C10)
    tekuca = None      # sekcija u koju upadaju pitanja

    def dodaj_sekciju(kod, naziv, dio, roditelj):
        s = {
            'kod': kod,
            'naziv': naziv,
            'dio': dio,
            'roditeljKod': roditelj,
            'redniBroj': len(sekcije) + 1,
            'pitanja': [],
        }
        sekcije.append(s)
        return s

    for r in range(PRVI_RED, POSLJEDNJI_RED_PITANJA + 1):
        a = ocisti(ws.cell(r, 1).value)
        b = ocisti(ws.cell(r, 2).value)
        c = ocisti(ws.cell(r, 3).value)
        d = ocisti(ws.cell(r, 4).value)
        g = ocisti(ws.cell(r, 7).value)

        if a in ('A', 'B', 'C'):
            dio = a
            oblast = None
            tekuca = None
            continue

        if d == 'Pitanje':          # red zaglavlja tabele
            continue

        m = RE_SEKCIJA.match(b)
        if m:
            kod = m.group(1)
            naziv = nazivi.get(kod, c) or c
            if dio == 'C':
                oblast = dodaj_sekciju(kod, naziv, dio, None)
                tekuca = oblast     # C9 i C10 nemaju podoblasti
            else:
                tekuca = dodaj_sekciju(kod, naziv, dio, None)
            continue

        m = RE_PODOBLAST.match(c)
        if m:
            kod = m.group(1)
            naziv = nazivi.get(kod, d) or d
            roditelj = oblast['kod'] if oblast else None
            tekuca = dodaj_sekciju(kod, naziv, dio, roditelj)
            continue

        if RE_BROJ_PITANJA.match(c) and d and tekuca is not None:
            tip, opcije = tip_odgovora(dio, d, 'E%d' % r, validacije)
            tekuca['pitanja'].append({
                'kod': '%s-%s' % (tekuca['kod'], c),
                'redniBroj': len(tekuca['pitanja']) + 1,
                'tekst': d,
                'uputstvo': g if dio in ('A', 'B') else '',
                'tipOdgovora': tip,
                'opcije': opcije,
            })

    # --- kontrola ---------------------------------------------------------
    for oznaka in ('A', 'B', 'C'):
        s = [x for x in sekcije if x['dio'] == oznaka]
        p = sum(len(x['pitanja']) for x in s)
        podoblasti = [x for x in s if x['roditeljKod']]
        print('  dio %s: %d sekcija (%d podoblasti), %d pitanja'
              % (oznaka, len(s), len(podoblasti), p))

    print()
    for s in sekcije:
        if s['dio'] == 'C' and s['roditeljKod'] is None:
            djeca = [x for x in sekcije if x['roditeljKod'] == s['kod']]
            ukupno = len(s['pitanja']) + sum(len(x['pitanja']) for x in djeca)
            print('  %-5s %4d pitanja  (%d podoblasti)  %s'
                  % (s['kod'], ukupno, len(djeca), s['naziv']))

    tipovi = {}
    for s in sekcije:
        for p in s['pitanja']:
            tipovi[p['tipOdgovora']] = tipovi.get(p['tipOdgovora'], 0) + 1
    print('\n  tipovi odgovora:', tipovi)

    kodovi = [p['kod'] for s in sekcije for p in s['pitanja']]
    assert len(kodovi) == len(set(kodovi)), 'duplirani kodovi pitanja'

    IZLAZ.parent.mkdir(parents=True, exist_ok=True)
    IZLAZ.write_text(
        json.dumps(sekcije, ensure_ascii=False, indent=2) + '\n',
        encoding='utf-8')

    print('\n  ukupno: %d sekcija, %d pitanja' % (len(sekcije), len(kodovi)))
    print('  zapisano: %s' % IZLAZ)


if __name__ == '__main__':
    main()
