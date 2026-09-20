# -*- coding: utf-8 -*-
"""
Izvlači katalog COBIT procesa i izjava iz izvornog Excel fajla u JSON.

Uzimaju se ISKLJUČIVO tekstovi izjava (kolone B i C). Popunjeni odgovori iz
izvornog fajla (kolone E–H) se NE prenose — riječ je o stvarnim odgovorima
stvarne institucije, a specifikacija zahtijeva da u aplikaciji budu samo
sintetički podaci.

Pokretanje:
    python alati/izvuci-cobit.py "putanja/do/2.1 COBIT.xlsx"

Izlaz: prisma/podaci/cobit.json
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

SHEETOVI = ['PO1', 'PO3', 'PO5', 'PO9', 'PO10', 'AI1', 'AI2', 'AI5', 'AI6',
            'DS1', 'DS4', 'DS5', 'DS10', 'DS11', 'ME1']

IZLAZ = Path(__file__).resolve().parent.parent / 'prisma' / 'podaci' / 'cobit.json'


def ocisti(tekst):
    return re.sub(r'\s+', ' ', str(tekst)).strip()


def nazivi_procesa(wb):
    """Nazivi se čitaju iz zbirnog sheeta: 'PO1 Strateško planiranje IS '."""
    nazivi = {}
    if 'Sheet1' not in wb.sheetnames:
        return nazivi
    ws = wb['Sheet1']
    for red in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        vrijednost = red[1].value
        if not vrijednost:
            continue
        pun = ocisti(vrijednost)
        podudaranje = re.match(r'^([A-Z]{2}\d+)\s+(.*)$', pun)
        if podudaranje:
            nazivi[podudaranje.group(1)] = podudaranje.group(2).strip()
    return nazivi


def izjave_sheeta(ws, kod_procesa):
    izjave = []
    brojac_po_nivou = {}

    for red in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3):
        a, b, c = red[0].value, red[1].value, red[2].value

        # Zbirni redovi ('Broj izjava' / 'Suma (nivo zrelosti)') se preskaču.
        if a is not None and 'broj izjava' in str(a).lower():
            continue
        if c is not None and str(c).strip().lower().startswith('suma'):
            continue
        if b is None or c is None:
            continue
        if not isinstance(b, (int, float)):
            continue

        nivo = int(b)
        if nivo < 0 or nivo > 5:
            continue

        tekst = ocisti(c)
        if not tekst:
            continue

        redni = brojac_po_nivou.get(nivo, 0) + 1
        brojac_po_nivou[nivo] = redni

        izjave.append({
            'kod': f'{kod_procesa}-N{nivo}-{redni}',
            'nivoZrelosti': nivo,
            'redniBroj': redni,
            'tekst': tekst,
        })

    return izjave


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    nazivi = nazivi_procesa(wb)

    procesi = []
    for redni, kod in enumerate(SHEETOVI, start=1):
        if kod not in wb.sheetnames:
            raise SystemExit(f'Nedostaje sheet: {kod}')

        izjave = izjave_sheeta(wb[kod], kod)
        procesi.append({
            'kod': kod,
            'naziv': nazivi.get(kod, kod),
            'domen': re.match(r'^[A-Z]+', kod).group(0),
            'redniBroj': redni,
            'izjave': izjave,
        })

        po_nivou = {}
        for i in izjave:
            po_nivou[i['nivoZrelosti']] = po_nivou.get(i['nivoZrelosti'], 0) + 1
        raspodjela = ' '.join(f'N{n}:{po_nivou.get(n, 0)}' for n in range(6))
        print(f'  {kod:<5} {len(izjave):>3} izjava   {raspodjela}   {nazivi.get(kod, "")}')

    IZLAZ.parent.mkdir(parents=True, exist_ok=True)
    IZLAZ.write_text(
        json.dumps(procesi, ensure_ascii=False, indent=2) + '\n',
        encoding='utf-8',
    )

    ukupno = sum(len(p['izjave']) for p in procesi)
    print(f'\n  ukupno: {len(procesi)} procesa, {ukupno} izjava')
    print(f'  zapisano: {IZLAZ}')


if __name__ == '__main__':
    main()
